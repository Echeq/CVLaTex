# -*- coding: utf-8 -*-
"""Word export through docx_gen: no LaTeX and no GUI.

The export mirrors the visual identity of each LaTeX template, so the layout
assertions below read the generated XML instead of counting paragraphs: a
section heading may be a body paragraph with a rule or a one-row plate/tab/bar
table, and the sidebar family keeps the whole body inside the body cell of the
page-wide header table. The helpers therefore walk the document tree in
document order, at every nesting depth the export uses (header table plus one
shaped table inside it).
"""

import base64
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import RGBColor  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

import docx_gen  # noqa: E402
import docx_profiles  # noqa: E402
import latex_gen  # noqa: E402

PREFS_PRUEBA = {
    "paper": "a4",
    "margins": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
    "font_size": 11,
}

#: Every template the Word export has to mirror.
ALL_TEMPLATES = tuple(latex_gen.TEMPLATES)

#: Decorative glyphs the playful heading puts before its title.
HEADING_MARKERS = "\u25cf\u25c6\u25aa\u25a0\u2022"

#: Profile bullet kind -> the geometric glyph the Word export draws.
BULLET_GLYPHS = {
    "dot": "\u2022",
    "square_navy": "\u25aa",
    "square_black": "\u25a0",
    "diamond_coral": "\u25c6",
}

# 1x1 PNG: enough for python-docx to embed and size the picture.
PNG_MINIMO = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQ"
    "DwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


# ---------------------------------------------------------------------------
# Document tree walking
# ---------------------------------------------------------------------------

def _cell_blocks(cell, depth):
    """Yield ``("paragraph"|"table", item, depth)`` inside a table cell."""
    for child in cell._tc.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, cell), depth
        elif child.tag == qn("w:tbl"):
            table = Table(child, cell)
            yield "table", table, depth + 1
            for item in _table_blocks(table, depth + 1):
                yield item


def _table_blocks(table, depth):
    """Yield the block items of every cell of a table in document order."""
    for row in table.rows:
        for cell in row.cells:
            for item in _cell_blocks(cell, depth):
                yield item


def _iter_blocks(doc):
    """Body blocks and table blocks in document order.

    The header is a body-level table, so it is yielded first; nested tables
    (the coloured band, the framed photo, the contact grid, a plate with its
    tail) follow inside their cell at depth 2.
    """
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, doc), 0
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            yield "table", table, 1
            for item in _table_blocks(table, 1):
                yield item


def _iter_paragraphs(doc):
    """Every paragraph in document order: body, header cells, nested tables."""
    for kind, item, _depth in _iter_blocks(doc):
        if kind == "paragraph":
            yield item


def _textos(doc):
    """All visible text, in document order."""
    return [paragraph.text for paragraph in _iter_paragraphs(doc)]


def _table_paragraphs(table):
    """Paragraphs of every cell of a table, nested tables included."""
    return [item for kind, item, _depth in _table_blocks(table, 1) if kind == "paragraph"]


def _tables_with_depth(doc):
    """``(table, depth)`` of every table; the header table is depth 1."""
    return [(item, depth) for kind, item, depth in _iter_blocks(doc) if kind == "table"]


def _ancestor_element(node, tag):
    """Closest ancestor (or self) of ``node`` with the Clark tag ``tag``."""
    element = node._p if isinstance(node, Paragraph) else node
    while element is not None:
        if element.tag == qn(tag):
            return element
        element = element.getparent()
    return None


def _ancestor_cell(paragraph):
    """``w:tc`` element that owns ``paragraph``, or ``None`` in the body flow."""
    return _ancestor_element(paragraph, "w:tc")


def _ancestor_table(paragraph):
    """``w:tbl`` element that owns ``paragraph``, or ``None`` in the body flow."""
    return _ancestor_element(paragraph, "w:tbl")


def _paragraph_after(doc, paragraph):
    """Paragraph that follows ``paragraph`` in document order.

    The stacked second rule of a two-tone heading is a sibling of a paragraph
    heading, but it follows the whole heading table for a table heading, so the
    flat document order is the only locator valid for both shapes.
    """
    paragraphs = list(_iter_paragraphs(doc))
    for index, candidate in enumerate(paragraphs):
        if candidate._p is paragraph._p:
            return paragraphs[index + 1] if index + 1 < len(paragraphs) else None
    return None


# ---------------------------------------------------------------------------
# XML readers (shading, borders, runs, geometry)
# ---------------------------------------------------------------------------

def _shading_fill(element):
    """``w:shd/@w:fill`` of an element that may carry ``w:shd``."""
    if element is None:
        return None
    shading = element.find(qn("w:shd"))
    if shading is None:
        return None
    return shading.get(qn("w:fill"))


def _paragraph_fill(paragraph):
    """Background fill of a paragraph, or ``None``."""
    return _shading_fill(paragraph._p.pPr)


def _cell_fill(cell):
    """Background fill of a table cell (``_Cell`` or raw ``w:tc``)."""
    tc = getattr(cell, "_tc", cell)
    if tc is None:
        return None
    return _shading_fill(tc.find(qn("w:tcPr")))


def _ancestor_cell_fill(paragraph):
    """Fill of the nearest ancestor cell of ``paragraph``."""
    return _cell_fill(_ancestor_cell(paragraph))


def _run_fill(run):
    """Run-box fill (``w:shd`` in ``w:rPr``), used by the playful badges."""
    return _shading_fill(run._element.find(qn("w:rPr")))


def _borders(element, tag="w:pBdr"):
    """Border dict of a ``w:pPr``/``w:tcPr``/``w:tblPr`` element.

    Keys are edge names (``top``, ``left``, ``bottom``...), values are dicts
    with ``val``, ``sz`` (eighths of a point) and ``color``.
    """
    if element is None:
        return {}
    container = element.find(qn(tag))
    if container is None:
        return {}
    out = {}
    for child in container:
        edge = child.tag.rsplit("}", 1)[-1]
        try:
            size = int(child.get(qn("w:sz")))
        except (TypeError, ValueError):
            size = None
        out[edge] = {
            "val": child.get(qn("w:val")),
            "sz": size,
            "color": child.get(qn("w:color")),
        }
    return out


def _paragraph_borders(paragraph):
    return _borders(paragraph._p.pPr)


def _table_borders(table):
    return _borders(table._tbl.tblPr, "w:tblBorders")


def _exact_line_pt(paragraph):
    """Exact rendered line height in points, or ``None`` when not exact."""
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None or spacing.get(qn("w:lineRule")) != "exact":
        return None
    try:
        return int(spacing.get(qn("w:line"))) / 20.0
    except (TypeError, ValueError):
        return None


def _text_runs(paragraph):
    """Runs that actually carry text."""
    return [run for run in paragraph.runs if run.text]


def _run_color(run):
    """Explicit run colour as RGBColor, or ``None`` when inherited."""
    color = run.font.color
    if color is None or color.type is None:
        return None
    return color.rgb


def _cell_width_cm(cell):
    """Exact ``w:tcW`` width of a cell in centimetres (dxa converted)."""
    tc = getattr(cell, "_tc", cell)
    tc_pr = tc.find(qn("w:tcPr"))
    width = tc_pr.find(qn("w:tcW")) if tc_pr is not None else None
    if width is None:
        return 0.0
    return int(width.get(qn("w:w"))) * 2.54 / 1440.0


def _grid_widths_cm(table):
    """Column widths of a fixed table, in centimetres."""
    grid = table._tbl.tblGrid
    if grid is None:
        return []
    return [int(node.get(qn("w:w"))) * 2.54 / 1440.0 for node in grid]


def _shaded_step_widths(cell):
    """Widths of the shaded step paragraphs of a blade cell.

    Each step is a shaded paragraph whose right indent grows, so the painted
    width shrinks from the cell edge. Only the shaded paragraphs are measured.
    """
    total = _cell_width_cm(cell)
    widths = []
    for paragraph in cell.paragraphs:
        if _paragraph_fill(paragraph) is None:
            continue
        fmt = paragraph.paragraph_format
        left = fmt.left_indent.cm if fmt.left_indent is not None else 0.0
        right = fmt.right_indent.cm if fmt.right_indent is not None else 0.0
        widths.append(total - left - right)
    return widths


def _strictly_decreasing(values):
    return len(values) >= 2 and all(a > b for a, b in zip(values, values[1:]))


# ---------------------------------------------------------------------------
# Header / heading locators
# ---------------------------------------------------------------------------

def _header_table(doc):
    """The borderless header table, always ``doc.tables[0]``."""
    assert doc.tables, "the header must be a table"
    return doc.tables[0]


def _header_name_paragraph(doc, name):
    """Name paragraph inside the header table (also when nested in a band)."""
    for paragraph in _table_paragraphs(_header_table(doc)):
        if paragraph.text.strip() == name:
            return paragraph
    raise AssertionError("No name paragraph %r in the header table" % name)


def _header_rule_paragraphs(doc):
    """Doc-level rule paragraphs between the header table and the body.

    The header rule(s) are the empty bordered paragraphs that follow the header
    table before the first text paragraph. The sidebar family draws its rule
    inside the body cell instead, so it returns an empty list there.
    """
    seen_table = False
    rules = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            if seen_table:
                break
            seen_table = True
            continue
        if not seen_table or child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc)
        if paragraph.text.strip():
            break
        if _paragraph_borders(paragraph):
            rules.append(paragraph)
    return rules


def _heading_text(text):
    """Heading text without the playful bullet glyph prefix."""
    return text.strip().lstrip(" \t" + HEADING_MARKERS)


def _find_heading(doc, title):
    """Locate a section heading by its exact title.

    Returns a dict with ``paragraph`` (the paragraph that holds the title),
    ``table``/``cell`` when the heading is a plate/tab/bar table, and
    ``depth`` (1 for a table heading). The body-flow rule headings have
    ``table`` set to ``None``.
    """
    for kind, item, depth in _iter_blocks(doc):
        if kind == "paragraph":
            if _heading_text(item.text) == title:
                return {
                    "paragraph": item, "table": None, "cell": None, "depth": depth,
                }
            continue
        # Heading tables are one-row fixed tables: only their direct cell
        # paragraphs hold the title (a nested shape never repeats it).
        if not item.rows:
            continue
        for cell in item.rows[0].cells:
            for paragraph in cell.paragraphs:
                if _heading_text(paragraph.text) == title:
                    return {
                        "paragraph": paragraph, "table": item, "cell": cell,
                        "depth": depth,
                    }
    raise AssertionError("No section heading %r" % title)


def _normal_rfonts(doc):
    """``w:rFonts`` attributes of the Normal style."""
    r_pr = doc.styles["Normal"].element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        return {}
    return {
        key: r_fonts.get(qn("w:" + key))
        for key in ("ascii", "hAnsi", "eastAsia", "cs")
    }


def _pg_borders(doc):
    """``w:pgBorders`` of the first section, or ``None``."""
    return doc.sections[0]._sectPr.find(qn("w:pgBorders"))


# ---------------------------------------------------------------------------
# Data builders
# ---------------------------------------------------------------------------

def _datos_docx():
    d = latex_gen.empty_data()
    d["cv_language"] = "es"
    d["name"] = "Ana García López"
    d["profession"] = "Ingeniera de software"
    d["email"] = "ana.garcia@example.com"
    d["summary"] = "Perfil profesional de prueba para la exportación Word."
    d["education"] = [{
        "degree_type": "bachelor", "degree": "Ingeniería Informática",
        "institution": "Universidad de Madrid", "start": "2012", "end": "2016",
        "grade": "8.5/10",
    }]
    d["work"] = [{
        "position": "Ingeniera backend", "company": "Fintech Nova", "start": "2020",
        "end": "Actualidad", "achievements": "Migramos a microservicios.",
    }]
    d["skills"] = [{"category": "Lenguajes", "values": "Python, Java"}]
    d["languages"] = [{"language": "Español", "level": "Nativo"}]
    d["references"] = [{"name": "Luis Pérez", "contact": "luis@example.com"}]
    return d


def _datos_con_plantilla(template):
    d = _datos_docx()
    d["template"] = template
    return d


def _datos_docx_con_bloque_educacion():
    d = _datos_docx()
    d["blocks"] = [
        {"id": "linea", "type": "separator", "style": "line",
         "position": "education"},
        {"id": "seccion", "type": "section", "title": "Mi sección",
         "content": "Primera línea del bloque.\nSegunda línea del bloque.",
         "position": "education"},
    ]
    return d


def _datos_docx_con_separadores():
    d = _datos_docx()
    d["blocks"] = [
        {"id": "linea", "type": "separator", "style": "line", "position": "end"},
        {"id": "banda", "type": "separator", "style": "band", "position": "end"},
    ]
    return d


def _exportar(datos, tmp_path, nombre):
    ruta = tmp_path / nombre
    docx_gen.export(datos, PREFS_PRUEBA, str(ruta))
    return Document(str(ruta))


def _textos_docx(datos, tmp_path, nombre):
    return _textos(_exportar(datos, tmp_path, nombre))


def _imagen_temporal(tmp_path, nombre="imagen.png"):
    ruta = tmp_path / nombre
    ruta.write_bytes(PNG_MINIMO)
    return str(ruta)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

def test_exportar_docx_incluye_nombre_y_secciones(tmp_path):
    doc = _exportar(_datos_docx(), tmp_path, "cv.docx")
    assert (tmp_path / "cv.docx").is_file()

    textos = _textos(doc)

    assert any("Ana García López" in t for t in textos)
    assert any("ana.garcia@example.com" in t for t in textos)

    for seccion in (
        "Educación", "Perfil", "Experiencia laboral",
        "Habilidades y certificados", "Idiomas", "Referencias",
    ):
        assert any(seccion in t for t in textos), seccion

    assert any("Universidad de Madrid" in t for t in textos)
    assert any("8.5/10" in t for t in textos)
    assert any("Python" in t for t in textos)
    assert any("Luis Pérez" in t for t in textos)


def test_exportar_docx_con_datos_vacios_no_falla(tmp_path):
    doc = _exportar(latex_gen.empty_data(), tmp_path, "vacio.docx")
    assert (tmp_path / "vacio.docx").is_file()
    assert _textos(doc) is not None


# ---------------------------------------------------------------------------
# Dynamic header: socials and extra fields
# ---------------------------------------------------------------------------

def test_exportar_docx_cabecera_socials(tmp_path):
    d = _datos_docx()
    d["socials"] = [
        {"network": "github", "label": "Ignorado", "value": "github.com/ana"},
        {"network": "linkedin", "label": "", "value": "linkedin.com/in/a"},
        {"network": "website", "label": "", "value": "ana.dev"},
        {"network": "other", "label": "Mastodon", "value": "mastodon/@ana"},
        {"network": "x", "label": "", "value": "   "},
    ]
    d["extra_fields"] = [{"label": "Nacimiento", "value": "12/05/1994"}]
    textos = _textos_docx(d, tmp_path, "cabecera_socials.docx")
    cabecera = " ".join(textos)

    assert "GitHub: github.com/ana" in cabecera
    assert "LinkedIn: linkedin.com/in/a" in cabecera
    assert "Web: ana.dev" in cabecera
    assert "Mastodon: mastodon/@ana" in cabecera
    assert "Nacimiento: 12/05/1994" in cabecera
    assert "Ignorado" not in cabecera
    # Personal data and extra fields, then the objective, then the socials.
    assert cabecera.index("Nacimiento:") < cabecera.index("Puesto al que aspiras:")
    assert cabecera.index("Puesto al que aspiras:") < cabecera.index("GitHub:")


def test_exportar_docx_objetivo_tras_datos_personales_y_sin_linea(tmp_path):
    """The contact block is a header-cell paragraph with no divider rule.

    The personal data, the extra fields and the objective share one paragraph
    (one line each with ``w:br``); the socials are a later, smaller and gray
    paragraph. Both live inside the header table, so the assertions read the
    recursive paragraph list instead of ``doc.paragraphs``.
    """
    d = _datos_docx()
    d["extra_fields"] = [{"label": "Nacimiento", "value": "12/05/1994"}]
    d["socials"] = [{"network": "github", "label": "", "value": "github.com/ana"}]
    doc = _exportar(d, tmp_path, "contactos.docx")

    parrafos = list(_iter_paragraphs(doc))
    contacto = next(p for p in parrafos if "ana.garcia@example.com" in p.text)
    texto = contacto.text
    # Personal data first, then the objective; socials live in a later line.
    assert texto.index("ana.garcia@example.com") < texto.index("Nacimiento:")
    assert texto.index("Nacimiento:") < texto.index("Puesto al que aspiras:")
    social = next(p for p in parrafos if "github.com/ana" in p.text)
    assert social is not contacto
    assert parrafos.index(social) > parrafos.index(contacto)
    assert "Puesto al que aspiras:" not in social.text

    # The social block is gray and carries no rule; its reduced size is pinned
    # by test_docx_sociales_mas_pequenos_que_contactos.
    github_run = next(run for run in _text_runs(social) if "github.com/ana" in run.text)
    assert _run_color(github_run) == RGBColor(0x6E, 0x76, 0x84)
    assert "bottom" not in _paragraph_borders(contacto)


def test_docx_sociales_mas_pequenos_que_contactos(tmp_path):
    """The social block is set smaller than the personal/objective block."""
    d = _datos_docx()
    d["socials"] = [{"network": "github", "label": "", "value": "github.com/ana"}]
    doc = _exportar(d, tmp_path, "social_size.docx")

    parrafos = list(_iter_paragraphs(doc))
    contacto = next(p for p in parrafos if "ana.garcia@example.com" in p.text)
    social = next(p for p in parrafos if "github.com/ana" in p.text)
    contacto_run = next(r for r in _text_runs(contacto) if r.text.strip())
    github_run = next(r for r in _text_runs(social) if "github.com/ana" in r.text)
    assert github_run.font.size.pt < contacto_run.font.size.pt
    assert github_run.font.size.pt < PREFS_PRUEBA["font_size"]


def test_exportar_docx_chinese_usa_un_gris_legible(tmp_path):
    """The light template gris is a bar colour, never a text colour.

    The Chinese template ``gris`` (#C4C8D4) paints the heading tail bar; the
    organization/degree text uses the readable default gray instead. Both facts
    are asserted, so a future change cannot silence the test by dropping either
    use of the colour.
    """
    doc = _exportar(_datos_con_plantilla("chinese"), tmp_path, "chinese.docx")

    heading = _find_heading(doc, "Perfil")
    assert heading["table"] is not None
    tail_cell = heading["table"].rows[0].cells[1]
    assert _cell_fill(tail_cell) == "C4C8D4"

    grado = next(p for p in _iter_paragraphs(doc) if "Grado Universitario" in p.text)
    assert [_run_color(run) for run in _text_runs(grado)] == [
        RGBColor(0x6E, 0x76, 0x84)
    ]
    # No run anywhere uses the light bar gray as text.
    colores_texto = {
        str(_run_color(run))
        for paragraph in _iter_paragraphs(doc)
        for run in _text_runs(paragraph)
    }
    assert "C4C8D4" not in colores_texto


# ---------------------------------------------------------------------------
# Free blocks
# ---------------------------------------------------------------------------

def test_exportar_docx_bloque_anclado_a_educacion(tmp_path):
    textos = _textos_docx(_datos_docx_con_bloque_educacion(), tmp_path, "bloques.docx")

    i_educacion = textos.index("Educación")
    i_bloque = textos.index("Mi sección")
    i_siguiente = textos.index("Experiencia laboral")
    assert i_educacion < i_bloque < i_siguiente
    assert any("Primera línea del bloque." in t for t in textos)
    assert any("Segunda línea del bloque." in t for t in textos)


def test_exportar_docx_bloque_seccion_vacio_se_omite(tmp_path):
    sin_bloque = _textos_docx(_datos_docx(), tmp_path, "sin_bloque.docx")

    d = _datos_docx()
    d["blocks"] = [
        {"id": "vacio", "type": "section", "title": "",
         "content": "", "position": "education"},
        {"id": "nulos", "type": "section", "title": None,
         "content": None, "position": "end"},
    ]
    con_vacios = _textos_docx(d, tmp_path, "con_vacios.docx")
    assert con_vacios == sin_bloque


def test_exportar_docx_bloques_listado(tmp_path):
    d = _datos_docx()
    d["blocks"] = [
        {"id": "l1", "type": "list", "mode": "columns", "title": "Enlaces",
         "columns": 2, "position": "end",
         "entries": [{"title": "Web", "detail": "example.com"}]},
        {"id": "l2", "type": "list", "mode": "period", "title": "Cronología",
         "columns": 1, "position": "end",
         "entries": [{"period": "2020", "description": "Hito uno."}]},
    ]
    textos = _textos_docx(d, tmp_path, "listado.docx")
    assert "Enlaces" in textos
    assert any("Web" in t and "example.com" in t for t in textos)
    assert "Cronología" in textos
    assert "2020" in textos
    assert any("Hito uno." in t for t in textos)


def test_exportar_docx_bloques_mal_formados_no_lanzan(tmp_path):
    d = _datos_docx()
    d["blocks"] = [
        None, 3, "texto", [], (),
        {},
        {"type": "other"},
        {"type": "section", "title": None, "content": None},
        {"type": "section", "title": 7, "content": 7},
        {"type": "separator", "style": "inexistente"},
        {"type": "separator", "style": None, "position": None},
        {"type": "separator", "position": "marte"},
    ]
    doc = _exportar(d, tmp_path, "mal_formados.docx")
    textos = _textos(doc)
    assert "7" in textos
    assert "Educación" in textos

    # A "blocks" value that is not a list does not break the export either.
    d["blocks"] = "no-es-lista"
    ruta2 = tmp_path / "no_lista.docx"
    docx_gen.export(d, PREFS_PRUEBA, str(ruta2))
    assert ruta2.is_file()


# ---------------------------------------------------------------------------
# Section order, hidden sections and custom titles
# ---------------------------------------------------------------------------

# Headings visible with _datos_docx() in the default order (only the sections
# that have content).
ETIQUETAS_ES = [
    "Educación", "Perfil", "Experiencia laboral", "Habilidades y certificados",
    "Idiomas", "Referencias",
]


def _indice(textos, aguja):
    """First position whose text contains ``aguja``; fails when absent."""
    for i, texto in enumerate(textos):
        if aguja in texto:
            return i
    raise AssertionError(f"No se encontro {aguja!r}")


def _secuencia_encabezados(textos, etiquetas):
    """Only the paragraphs that are exactly a section heading."""
    conocidas = set(etiquetas)
    return [t for t in textos if t in conocidas]


def test_exportar_docx_orden_por_defecto_no_cambia(tmp_path):
    textos = _textos_docx(_datos_docx(), tmp_path, "defecto.docx")
    assert _secuencia_encabezados(textos, ETIQUETAS_ES) == ETIQUETAS_ES


def test_exportar_docx_orden_secciones_invertido(tmp_path):
    d = _datos_docx()
    d["section_order"] = list(reversed(latex_gen.DEFAULT_SECTION_ORDER))
    textos = _textos_docx(d, tmp_path, "orden_inverso.docx")

    assert _secuencia_encabezados(textos, ETIQUETAS_ES) == [
        "Referencias", "Idiomas", "Habilidades y certificados",
        "Experiencia laboral", "Perfil", "Educación",
    ]

    # The entries travel with their section, not only the headings.
    indices = [
        _indice(textos, aguja) for aguja in (
            "Luis Pérez",                       # references
            "Español",                          # languages
            "Python, Java",                     # skills
            "Ingeniera backend",                # work
            "Perfil profesional de prueba",     # summary
            "Universidad de Madrid",            # education
        )
    ]
    assert indices == sorted(indices)


def test_exportar_docx_seccion_oculta_sin_encabezado_ni_entradas(tmp_path):
    d = _datos_docx()
    d["hidden_sections"] = ["education"]
    textos = _textos_docx(d, tmp_path, "oculta.docx")

    assert "Educación" not in textos
    assert not any("Universidad de Madrid" in t for t in textos)
    assert not any("Ingeniería Informática" in t for t in textos)
    assert not any("8.5/10" in t for t in textos)

    assert _secuencia_encabezados(textos, ETIQUETAS_ES) == [
        "Perfil", "Experiencia laboral", "Habilidades y certificados",
        "Idiomas", "Referencias",
    ]


def test_exportar_docx_titulos_seccion_sustituyen_etiquetas(tmp_path):
    d = _datos_docx()
    d["section_titles"] = {
        "education": "Formación académica",
        "work": "Trayectoria profesional",
        "skills": "Competencias",
    }
    textos = _textos_docx(d, tmp_path, "titulos.docx")

    esperados = [
        "Formación académica", "Perfil", "Trayectoria profesional",
        "Competencias", "Idiomas", "Referencias",
    ]
    assert _secuencia_encabezados(textos, esperados) == esperados
    for etiqueta in ("Educación", "Experiencia laboral", "Habilidades"):
        assert etiqueta not in textos
    assert any("Universidad de Madrid" in t for t in textos)
    assert any("Ingeniera backend" in t for t in textos)
    assert any("Python" in t for t in textos)


def test_exportar_docx_bloque_anclado_a_seccion_oculta_conserva_posicion(tmp_path):
    d = _datos_docx()
    d["hidden_sections"] = ["work"]
    d["blocks"] = [{
        "id": "nota", "type": "section", "title": "Nota de trayectoria",
        "content": "Bloque anclado a una seccion oculta.",
        "position": "work",
    }]
    textos = _textos_docx(d, tmp_path, "oculta_con_bloque.docx")

    assert "Experiencia laboral" not in textos
    assert not any("Ingeniera backend" in t for t in textos)
    assert not any("Migramos a microservicios." in t for t in textos)

    i_educacion = textos.index("Educación")
    i_bloque = textos.index("Nota de trayectoria")
    i_habilidades = textos.index("Habilidades y certificados")
    assert i_educacion < i_bloque < i_habilidades
    assert i_educacion < _indice(textos, "Bloque anclado a una seccion oculta.") < i_habilidades


def test_exportar_docx_bloque_seccion_lleva_regla_de_color(tmp_path):
    """A free section block gets the template heading shape, not a bar.

    The classic heading is a coloured title over a 0.75pt accent rule (the old
    export filled a single generic heading paragraph); the block still sits
    between its anchor and the next section.
    """
    doc = _exportar(_datos_docx_con_bloque_educacion(), tmp_path, "bloque_barra.docx")

    heading = _find_heading(doc, "Mi sección")
    assert heading["table"] is None
    bottom = _paragraph_borders(heading["paragraph"])["bottom"]
    assert bottom["color"] == docx_gen.COLOR_ACCENT
    assert bottom["sz"] == 6  # 0.75pt stroke

    textos = _textos(doc)
    assert textos.index("Educación") < textos.index("Mi sección") < textos.index("Experiencia laboral")
    assert textos.index("Mi sección") < textos.index("Primera línea del bloque.")


# ---------------------------------------------------------------------------
# Visual header: borderless table, photo/logo and long names
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("template", ALL_TEMPLATES)
def test_exportar_docx_cabecera_en_tabla_sin_bordes(tmp_path, template):
    d = _datos_con_plantilla(template)
    doc = _exportar(d, tmp_path, "cabecera_%s.docx" % template)

    tabla = _header_table(doc)
    bordes = _table_borders(tabla)
    assert set(bordes) == {
        "top", "left", "bottom", "right", "insideH", "insideV",
    }
    assert all(borde["val"] == "nil" for borde in bordes.values())
    assert len(tabla.rows) == 1
    assert len(tabla.columns) >= 1
    assert any(
        paragraph.text.strip() == d["name"]
        for paragraph in _table_paragraphs(tabla)
    )


def test_exportar_docx_incrusta_foto_y_logo_en_dos_columnas(tmp_path):
    d = _datos_docx()
    d["photo"] = _imagen_temporal(tmp_path, "foto.png")
    d["school_logo"] = _imagen_temporal(tmp_path, "logo.png")
    doc = _exportar(d, tmp_path, "con_imagenes.docx")

    assert len(doc.tables[0].columns) == 2
    assert len(doc.inline_shapes) == 2
    # The images are right-aligned in their cell (the logo paragraph is the
    # direct child; the framed photo lives in a nested table).
    derecha = doc.tables[0].rows[0].cells[1]
    assert any(p.alignment == WD_ALIGN_PARAGRAPH.RIGHT for p in derecha.paragraphs)


def test_exportar_docx_rutas_de_imagen_invalidas_se_ignoran(tmp_path):
    d = _datos_docx()
    d["photo"] = str(tmp_path / "no_existe.png")
    d["school_logo"] = str(tmp_path / "no_existe.jpg")
    doc = _exportar(d, tmp_path, "sin_imagenes.docx")

    assert len(doc.tables[0].columns) == 1
    assert len(doc.inline_shapes) == 0

    # An unsupported extension is never embedded either.
    d["photo"] = _imagen_temporal(tmp_path, "foto.gif")
    ruta2 = tmp_path / "extension_invalida.docx"
    docx_gen.export(d, PREFS_PRUEBA, str(ruta2))
    assert len(Document(str(ruta2)).inline_shapes) == 0


def test_exportar_docx_nombre_vacio_usa_el_titulo_corto_una_vez(tmp_path):
    d = _datos_docx()
    d["name"] = "   "
    textos = _textos_docx(d, tmp_path, "sin_nombre.docx")
    assert textos.count("CURRÍCULUM") == 1


def test_exportar_docx_nombre_largo_no_falla_y_reduce_el_tamano(tmp_path):
    d = _datos_docx()
    d["name"] = (
        "María de los Ángeles Fernández de la Vega y Rodríguez de los Ríos "
        "del Valle"
    )
    doc = _exportar(d, tmp_path, "nombre_largo.docx")

    assert any(d["name"] in t for t in _textos(doc))
    nombre = _header_name_paragraph(doc, d["name"])
    tamanos = [run.font.size.pt for run in _text_runs(nombre) if run.font.size]
    base = docx_profiles.PROFILES["classic"].name_pt
    assert tamanos and max(tamanos) < base


def test_exportar_docx_plantilla_hostil_no_falla(tmp_path):
    for i, plantilla in enumerate(({"x": 1}, 7, None, "", "marte")):
        d = _datos_docx()
        d["template"] = plantilla
        ruta = tmp_path / ("plantilla_%d.docx" % i)
        docx_gen.export(d, PREFS_PRUEBA, str(ruta))
        assert ruta.is_file()
        assert any("Ana García López" in t for t in _textos(Document(str(ruta))))


# ---------------------------------------------------------------------------
# Per-template signature
# ---------------------------------------------------------------------------

def _assert_normal_fonts(doc, template):
    """Every template writes explicit ascii/hAnsi/eastAsia fonts on Normal."""
    fonts = _normal_rfonts(doc)
    assert fonts.get("eastAsia") == docx_gen.EAST_ASIA_FONT
    assert fonts.get("ascii") == fonts.get("hAnsi")
    profile = docx_profiles.PROFILES[template]
    expected = docx_gen.SANS_FONT if profile.font_family == "sans" else docx_gen.SERIF_FONT
    assert fonts.get("ascii") == expected


def _assert_template_signature(doc, template, d, palette):
    """Assert the identity element that makes ``template`` recognizable."""
    name = _header_name_paragraph(doc, d["name"])
    name_colors = [_run_color(run) for run in _text_runs(name)]

    if template == "classic":
        # Black name over a 0.75pt accent header rule.
        assert name_colors == [RGBColor(0, 0, 0)]
        rules = _header_rule_paragraphs(doc)
        assert len(rules) == 1
        bottom = _paragraph_borders(rules[0])["bottom"]
        assert bottom["color"] == palette.accent_hex
        assert bottom["sz"] == 6

    elif template == "harvard":
        # Centred small-caps name over a gray double hairline.
        assert name.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert _text_runs(name)[0].font.small_caps is True
        rules = _header_rule_paragraphs(doc)
        assert len(rules) == 2
        bottoms = [_paragraph_borders(p)["bottom"] for p in rules]
        assert [b["color"] for b in bottoms] == [palette.gray_hex, palette.gray_hex]
        # Word stores the 0.5pt and 0.25pt strokes as 4 and 2 eighths.
        assert [b["sz"] for b in bottoms] == [4, 2]

    elif template == "modern":
        # Big accent name and a 2.5pt accent rule.
        assert name_colors == [palette.accent]
        rules = _header_rule_paragraphs(doc)
        assert len(rules) == 1
        bottom = _paragraph_borders(rules[0])["bottom"]
        assert bottom["color"] == palette.accent_hex
        assert bottom["sz"] == 20

    elif template == "creative":
        # White name on the accent band; teal + yellow stacked heading rules.
        assert _ancestor_cell_fill(name) == palette.accent_hex
        assert name_colors == [RGBColor(0xFF, 0xFF, 0xFF)]
        heading = _find_heading(doc, "Perfil")
        first = _paragraph_borders(heading["paragraph"])["bottom"]
        assert first["color"] == palette.secondary_hex
        second = _paragraph_after(doc, heading["paragraph"])
        assert second is not None and not second.text
        assert _paragraph_borders(second)["bottom"]["color"] == palette.accent3_hex
        assert palette.accent3_hex == "F4B440"

    elif template == "chinese":
        # Accent plate behind the white name; heading plate + light tail bar.
        # The profile names an accent2 tail, but this template has no acento2
        # (it would resolve to the accent and merge with the plate), so the tail
        # uses the template gris as the LaTeX template does: the identity is a
        # solid accent plate next to a distinct, much lighter tail bar.
        assert _ancestor_cell_fill(name) == palette.accent_hex
        assert name_colors == [RGBColor(0xFF, 0xFF, 0xFF)]
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is not None
        plate, tail = heading["table"].rows[0].cells[0], heading["table"].rows[0].cells[1]
        assert _cell_fill(plate) == palette.accent_hex
        assert _cell_fill(tail) == palette.template_gray_hex
        assert _cell_fill(tail) != _cell_fill(plate)

    elif template == "tabbed":
        # Page frame (sz 14) and a black tab heading with a white title.
        pg = _pg_borders(doc)
        assert pg is not None
        assert pg.get(qn("w:offsetFrom")) == "page"
        assert {child.get(qn("w:sz")) for child in pg} == {"14"}
        assert name_colors == [RGBColor(0, 0, 0)]
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is not None
        cells = heading["table"].rows[0].cells
        assert _cell_fill(cells[0]) == palette.accent_hex
        assert [_run_color(run) for run in _text_runs(heading["paragraph"])] == [
            RGBColor(0xFF, 0xFF, 0xFF)
        ]
        rule = _paragraph_borders(cells[-1].paragraphs[0])["bottom"]
        assert rule["color"] == palette.secondary_hex

    elif template == "ribbon":
        # Gold tick on the accent bar and the gold rule under it.
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is not None
        cell = heading["table"].rows[0].cells[0]
        assert _cell_fill(cell) == palette.accent_hex
        tint = _paragraph_borders(heading["paragraph"])["left"]
        assert tint["color"] == palette.secondary_hex
        assert tint["val"] == "single"
        rule = _paragraph_borders(_paragraph_after(doc, heading["paragraph"]))["bottom"]
        assert rule["color"] == palette.secondary_hex
        assert rule["sz"] == 10  # 1.2pt gold rule
        # The ribbon wing itself keeps both accent and gold bars.
        fills = {_paragraph_fill(p) for p in _table_paragraphs(_header_table(doc))}
        assert {palette.accent_hex, palette.secondary_hex} <= fills

    elif template == "navy":
        # Two stacked hairlines: 0.6pt accent + 0.3pt secondary.
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is None
        first = _paragraph_borders(heading["paragraph"])["bottom"]
        assert first["color"] == palette.accent_hex
        assert first["sz"] == 5  # 0.6pt -> 5/8pt in Word
        second = _paragraph_after(doc, heading["paragraph"])
        assert second is not None and not second.text
        second_rule = _paragraph_borders(second)["bottom"]
        assert second_rule["color"] == palette.secondary_hex
        assert second_rule["sz"] == 2  # 0.3pt -> 2/8pt in Word

    elif template == "playful":
        # Sans-serif Normal, purple round badge and coral diamond bullets.
        assert doc.styles["Normal"].font.name == docx_gen.SANS_FONT
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is None
        badge = _text_runs(heading["paragraph"])[0]
        assert badge.text.strip() == "\u25cf"
        assert _run_color(badge) == palette.accent
        badges = [
            run
            for paragraph in _table_paragraphs(_header_table(doc))
            for run in paragraph.runs
            if _run_fill(run) == palette.accent_hex
        ]
        assert badges
        assert all(_run_color(run) == RGBColor(0xFF, 0xFF, 0xFF) for run in badges)
        logro = next(p for p in _iter_paragraphs(doc) if "Migramos" in p.text)
        marker = _text_runs(logro)[0]
        assert marker.text.startswith("\u25c6")
        assert _run_color(marker) == palette.secondary

    elif template == "sword":
        # Accent plate + accent2 full-width bar + stair blade of 12 steps.
        assert _ancestor_cell_fill(name) == palette.accent_hex
        assert name_colors == [RGBColor(0xFF, 0xFF, 0xFF)]
        heading = _find_heading(doc, "Perfil")
        assert heading["table"] is not None
        cells = heading["table"].rows[0].cells
        assert _cell_fill(cells[0]) == palette.accent_hex
        assert _cell_fill(cells[2]) == palette.secondary_hex
        assert _cell_width_cm(cells[2]) > _cell_width_cm(cells[0])
        assert abs(sum(_cell_width_cm(cell) for cell in cells) - 17.0) < 0.05
        steps = _shaded_step_widths(cells[1])
        assert len(steps) >= 6
        assert _strictly_decreasing(steps)
        # The header banner has its own blade tip: two blades in this document.
        blades = [
            widths
            for table, _depth in _tables_with_depth(doc)
            for row in table.rows
            for cell in row.cells
            for widths in [_shaded_step_widths(cell)]
            if len(widths) >= 6
        ]
        assert len(blades) >= 2
        assert all(_strictly_decreasing(widths) for widths in blades)

    elif template in ("sidebar", "sidebar-right"):
        table = _header_table(doc)
        indent = table._tbl.tblPr.find(qn("w:tblInd"))
        assert indent is not None
        assert int(indent.get(qn("w:w"))) < 0
        mirrored = template.endswith("-right")
        bar_index, body_index = (3, 1) if mirrored else (0, 2)
        cells = table.rows[0].cells
        bar = cells[bar_index]
        assert _cell_fill(bar) == palette.accent_hex
        contacts = [p for p in bar.paragraphs if "ana.garcia@example.com" in p.text]
        assert contacts
        email = next(r for r in _text_runs(contacts[0]) if "ana.garcia@example.com" in r.text)
        assert _run_color(email) == RGBColor(0xFF, 0xFF, 0xFF)
        # The body (name, headings, entries) lives inside the body cell of the
        # same page-wide table, not in the normal flow nor in the bar cell.
        heading = _find_heading(doc, "Perfil")
        body_cell = _ancestor_cell(heading["paragraph"])
        assert body_cell is not None
        assert body_cell is not _ancestor_cell(contacts[0])
        assert body_cell is cells[body_index]._tc
        assert _ancestor_table(heading["paragraph"]) is table._tbl
        assert _ancestor_cell(name) is cells[body_index]._tc
        assert cells[bar_index]._tc is not cells[body_index]._tc

    else:  # pragma: no cover - the parametrisation covers every template
        raise AssertionError("no signature for %r" % template)


@pytest.mark.parametrize("template", ALL_TEMPLATES)
def test_docx_firma_de_plantilla(tmp_path, template):
    d = _datos_con_plantilla(template)
    doc = _exportar(d, tmp_path, "firma_%s.docx" % template)
    palette = docx_gen._build_palette(d, PREFS_PRUEBA)

    _assert_normal_fonts(doc, template)
    _assert_template_signature(doc, template, d, palette)


def test_docx_arbol_de_tablas_como_mucho_dos_niveles(tmp_path):
    """Every template keeps the header table plus one nested shape (depth 2)."""
    for template in ALL_TEMPLATES:
        doc = _exportar(
            _datos_con_plantilla(template), tmp_path, "prof_%s.docx" % template
        )
        depths = [depth for _table, depth in _tables_with_depth(doc)]
        assert depths, template
        assert max(depths) <= 2, template
        assert min(depths) == 1, template


# ---------------------------------------------------------------------------
# Entry lines and columns
# ---------------------------------------------------------------------------

def test_docx_entrada_clasica_fecha_con_tab_derecha(tmp_path):
    doc = _exportar(_datos_docx(), tmp_path, "tab_fecha.docx")
    entrada = next(p for p in _iter_paragraphs(doc) if "Ingeniera backend" in p.text)
    stops = list(entrada.paragraph_format.tab_stops)
    assert len(stops) == 1
    assert stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
    assert abs(stops[0].position.cm - 17.0) < 0.05


def test_docx_entrada_cn_tres_runs_y_columnas_ignoradas(tmp_path):
    """Chinese entries are one three-run line and ignore ``columns``."""
    d = _datos_con_plantilla("chinese")
    d["columns"] = {"work": 3, "education": 3}
    doc = _exportar(d, tmp_path, "cn_columnas.docx")

    entrada = next(p for p in _iter_paragraphs(doc) if "Fintech Nova" in p.text)
    # It stays in the normal flow: no per-entry column grid is built.
    assert _ancestor_table(entrada) is None
    texts = [run.text for run in _text_runs(entrada)]
    assert len(texts) == 3
    assert texts[0].startswith("2020")
    assert "Ingeniera backend" in texts[1]
    assert "Fintech Nova" in texts[2]
    assert len(entrada.text.splitlines()) == 1


def test_docx_columnas_2_y_3_respetadas_en_classic(tmp_path):
    d = _datos_docx()
    d["columns"] = {"work": 2, "education": 3}
    d["work"] = [
        {"position": "Ingeniera backend", "company": "Fintech Nova",
         "start": "2020", "end": "Actualidad"},
        {"position": "Desarrolladora", "company": "Acme", "start": "2016",
         "end": "2020"},
    ]
    d["education"] = [
        {"degree_type": "bachelor", "degree": "Informática",
         "institution": "Universidad de Madrid", "start": "2012", "end": "2016"},
        {"degree_type": "master", "degree": "Software",
         "institution": "Universidad de Barcelona", "start": "2016", "end": "2018"},
        {"degree_type": "course", "degree": "Cloud",
         "institution": "Escuela Cloud", "start": "2019"},
    ]
    doc = _exportar(d, tmp_path, "columnas.docx")

    texto_ancho = 17.0  # A4 minus the two 2cm margins

    work_table = next(
        table for table in doc.tables
        if any("Desarrolladora" in cell.text for row in table.rows for cell in row.cells)
    )
    assert len(work_table.columns) == 2
    assert len(work_table.rows) == 1
    assert abs(sum(_grid_widths_cm(work_table)) - texto_ancho) < 0.05
    work_cells = [cell.text for cell in work_table.rows[0].cells]
    assert any("Ingeniera backend" in t for t in work_cells)
    assert any("Desarrolladora" in t for t in work_cells)
    assert len({t for t in work_cells if t.strip()}) == 2

    edu_table = next(
        table for table in doc.tables
        if any("Escuela Cloud" in cell.text for row in table.rows for cell in row.cells)
    )
    assert len(edu_table.columns) == 3
    assert len(edu_table.rows) == 1
    assert abs(sum(_grid_widths_cm(edu_table)) - texto_ancho) < 0.05
    edu_cells = [cell.text for cell in edu_table.rows[0].cells]
    for institucion in (
        "Universidad de Madrid", "Universidad de Barcelona", "Escuela Cloud",
    ):
        assert sum(institucion in t for t in edu_cells) == 1


# ---------------------------------------------------------------------------
# Separators and bullets
# ---------------------------------------------------------------------------

def test_docx_separadores_de_bloques_finos(tmp_path):
    """Classic line + band separators are thin rules, not a filled band."""
    doc = _exportar(_datos_docx_con_separadores(), tmp_path, "separadores.docx")

    parrafos = list(_iter_paragraphs(doc))
    linea, banda = parrafos[-2], parrafos[-1]
    assert not linea.text and not banda.text

    borde = _paragraph_borders(linea)["bottom"]
    assert borde["color"] == docx_gen.COLOR_ACCENT
    assert borde["sz"] <= 8  # 0.6pt hairline
    assert _exact_line_pt(linea) is not None and _exact_line_pt(linea) < 2.0

    assert _paragraph_fill(banda) == docx_gen.COLOR_ACCENT
    altura = _exact_line_pt(banda)
    assert altura is not None
    assert 1.0 <= altura < 10.0  # the 3.5pt band, never an 11pt filled block


@pytest.mark.parametrize("template", ALL_TEMPLATES)
def test_docx_ninguna_banda_rellena_alta(tmp_path, template):
    """No template paints a separator as a tall filled paragraph."""
    d = _datos_docx_con_separadores()
    d["template"] = template
    doc = _exportar(d, tmp_path, "banda_%s.docx" % template)

    for paragraph in _iter_paragraphs(doc):
        altura = _exact_line_pt(paragraph)
        if _paragraph_fill(paragraph) and altura is not None:
            assert altura < 10.0, (template, paragraph.text, altura)


@pytest.mark.parametrize("template", ALL_TEMPLATES)
def test_docx_vineta_de_perfil(tmp_path, template):
    """The achievements use the template's geometric bullet, hanging indent."""
    d = _datos_con_plantilla(template)
    doc = _exportar(d, tmp_path, "vineta_%s.docx" % template)

    kind = docx_profiles.PROFILES[template].bullet_kind
    glyph = BULLET_GLYPHS[kind]
    logro = next(p for p in _iter_paragraphs(doc) if "Migramos" in p.text)
    marker = _text_runs(logro)[0]
    assert marker.text.startswith(glyph)
    assert logro.paragraph_format.left_indent is not None
    assert logro.paragraph_format.left_indent.cm > 0
