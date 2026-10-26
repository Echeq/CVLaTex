# -*- coding: utf-8 -*-
"""Paleta de colores del documento: normalizacion, LaTeX, API y Word.

Sin GUI ni pywebview. La compilacion PDF real lleva skipif cuando no hay
pdflatex (mismo patron que ``tests/test_roundtrip_tex.py``). Los valores de
los colores de plantilla se contrastan contra los ``templates/*.tex`` reales.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import RGBColor  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

import docx_gen  # noqa: E402
import latex_engine  # noqa: E402
import latex_gen  # noqa: E402
import prefs  # noqa: E402
import tex_datos  # noqa: E402
from main import Api  # noqa: E402

# Misma forma que tests/test_esquema.py: sin la clave "colors".
PREFS_PRUEBA = {
    "paper": "a4",
    "margins": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
    "font_size": 11,
}

COLORES_VALIDOS = {
    "primary": "#12AB34",
    "secondary": "#00AA88",
    "text": "#112233",
    "gray": "#A1B2C3",
}

MOTOR_PDF = latex_engine.find_pdflatex()

ETIQUETAS_ES = ("Perfil", "Educación", "Experiencia laboral")


class _Explosivo:
    """Tipo hostil: convertirlo a texto o compararlo revienta."""

    def __str__(self):
        raise RuntimeError("boom")

    def __eq__(self, otro):
        raise RuntimeError("boom")

    __hash__ = None


class _DictExplosivo(dict):
    """Diccionario hostil: cualquier get() revienta."""

    def get(self, clave, defecto=None):
        raise RuntimeError("boom")


def _prefs_con(colores):
    p = dict(PREFS_PRUEBA)
    p["colors"] = colores
    return p


def _datos_prueba():
    d = latex_gen.empty_data()
    d["cv_language"] = "es"
    d["summary"] = "Perfil con colores."
    d["work"] = [{"position": "Desarrolladora", "company": "Acme"}]
    return d


def _tex_de_plantilla(pid):
    ruta = os.path.join(latex_gen.BASE_DIR, "templates", latex_gen.TEMPLATE_FILES[pid])
    with open(ruta, "r", encoding="utf-8") as f:
        return f.read()


# ---------------------------------------------------------------------------
# normalizar_hex
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "valor, esperado",
    [
        ("#12AB34", (0x12, 0xAB, 0x34)),
        ("#12ab34", (0x12, 0xAB, 0x34)),
        ("  #12AB34  ", (0x12, 0xAB, 0x34)),
        ("#abc", (0xAA, 0xBB, 0xCC)),
        ("#AbC", (0xAA, 0xBB, 0xCC)),
        ("  #aBc  ", (0xAA, 0xBB, 0xCC)),
        ("#000000", (0, 0, 0)),
        ("#FFFFFF", (255, 255, 255)),
    ],
)
def test_normalizar_hex_acepta_formas_validas(valor, esperado):
    assert latex_gen.parse_hex_color(valor) == esperado


@pytest.mark.parametrize(
    "valor",
    [
        "",
        "   ",
        None,
        "red",
        "#12",
        "#12345",
        "#GGGGGG",
        "#12345Z",
        123,
        ["#FFFFFF"],
        {"x": 1},
        b"#FFFFFF",
    ],
)
def test_normalizar_hex_rechaza_invalidos(valor):
    assert latex_gen.parse_hex_color(valor) is None


# ---------------------------------------------------------------------------
# colores_de_plantilla sobre los templates reales
# ---------------------------------------------------------------------------

def test_colores_de_plantilla_classic():
    assert latex_gen.template_colors(_tex_de_plantilla("classic")) == {
        "accent": "#233A5E",
        "gris": "#6E7684",
    }


def test_colores_de_plantilla_creative():
    colores = latex_gen.template_colors(_tex_de_plantilla("creative"))
    assert colores["accent"] == "#E05244"
    assert colores["acento2"] == "#00897B"
    # El template real define acento3 como RGB{244,180,64} (#F4B440).
    assert colores["acento3"] == "#F4B440"


def test_colores_de_plantilla_chinese():
    assert latex_gen.template_colors(_tex_de_plantilla("chinese"))["accent"] == "#1F347D"


@pytest.mark.parametrize("tex", ["", None, 7, b"%", _Explosivo()])
def test_colores_de_plantilla_no_lanza(tex):
    assert latex_gen.template_colors(tex) == {}


# ---------------------------------------------------------------------------
# build_tex: sustituciones de la paleta
# ---------------------------------------------------------------------------

def test_build_tex_primary_reemplaza_accent():
    tex = latex_gen.build_tex(_datos_prueba(), _prefs_con({"primary": "#123456"}))
    assert "\\definecolor{accent}{RGB}{18,52,86}" in tex
    assert "35,58,94" not in tex


def test_build_tex_secondary_reemplaza_acento2_en_creative():
    d = _datos_prueba()
    d["template"] = "creative"
    tex = latex_gen.build_tex(d, _prefs_con({"secondary": "#00AA88"}))
    assert "\\definecolor{acento2}{RGB}{0,170,136}" in tex
    assert tex.count("\\definecolor{acento2}") == 1


def test_build_tex_secondary_en_classic_inyecta_acento2():
    # Comportamiento congelado de t1: classic no define acento2, asi que la
    # definicion se inyecta antes de \begin{document} en vez de reemplazar
    # nada; las definiciones existentes (accent y gris) quedan intactas.
    tex = latex_gen.build_tex(_datos_prueba(), _prefs_con({"secondary": "#00AA88"}))
    inyectado = "\\definecolor{acento2}{RGB}{0,170,136}"
    assert inyectado in tex
    assert tex.index(inyectado) < tex.index("\\begin{document}")
    assert "\\definecolor{accent}{RGB}{35,58,94}" in tex
    assert "\\definecolor{gris}{RGB}{110,118,132}" in tex


def test_build_tex_gray_reemplaza_gris():
    tex = latex_gen.build_tex(_datos_prueba(), _prefs_con({"gray": "#A1B2C3"}))
    assert "\\definecolor{gris}{RGB}{161,178,195}" in tex
    assert "110,118,132" not in tex


def test_build_tex_text_inyecta_cvtexto_antes_y_color_despues():
    tex = latex_gen.build_tex(_datos_prueba(), _prefs_con({"text": "#112233"}))
    definicion = tex.index("\\definecolor{cvtexto}{RGB}{17,34,51}")
    begin = tex.index("\\begin{document}")
    activacion = tex.index("\\color{cvtexto}", begin)
    assert definicion < begin < activacion
    assert tex.count("\\definecolor{cvtexto}") == 1
    assert tex.count("\\color{cvtexto}") == 1


def test_build_tex_roundtrip_con_paleta_completa():
    d = _datos_prueba()
    tex = latex_gen.build_tex(d, _prefs_con(COLORES_VALIDOS))
    for esperado in (
        "\\definecolor{accent}{RGB}{18,171,52}",
        "\\definecolor{acento2}{RGB}{0,170,136}",
        "\\definecolor{gris}{RGB}{161,178,195}",
        "\\definecolor{cvtexto}{RGB}{17,34,51}",
    ):
        assert esperado in tex
    assert tex_datos.extract(tex) == d


# ---------------------------------------------------------------------------
# build_tex: identidad byte a byte sin colores validos
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "colores",
    [
        None,  # prefs sin la clave "colors"
        "no-es-dict",
        {"primary": "", "secondary": "", "text": "", "gray": ""},
        {"primary": "red", "secondary": "#12", "text": "#GGGGGG", "gray": "  "},
        {"primary": 123, "secondary": [], "text": {"x": 1}, "gray": None},
        {"primary": "#", "secondary": "#12345", "text": "#12345Z", "gray": True},
    ],
)
def test_build_tex_sin_colores_validos_no_cambia_la_salida(colores):
    d = _datos_prueba()
    p = dict(PREFS_PRUEBA)
    if colores is not None:
        p["colors"] = colores
    assert latex_gen.build_tex(d, p) == latex_gen.build_tex(d, PREFS_PRUEBA)


def test_build_tex_colores_hostiles_no_cambian_la_salida():
    d = _datos_prueba()
    base = latex_gen.build_tex(d, PREFS_PRUEBA)
    for colores in (_Explosivo(), _DictExplosivo(), _DictExplosivo(primary="#123456")):
        assert latex_gen.build_tex(d, _prefs_con(colores)) == base


# ---------------------------------------------------------------------------
# Api.template_colors
# ---------------------------------------------------------------------------

def test_api_template_colors_creative():
    r = Api(None).template_colors("creative")
    assert r["ok"] is True
    assert r["colores"]["accent"] == "#E05244"
    assert r["colores"]["acento2"] == "#00897B"


@pytest.mark.parametrize("plantilla", [None, "marte", "", 123, []])
def test_api_template_colors_desconocida_cae_a_classic(plantilla):
    r = Api(None).template_colors(plantilla)
    assert r["ok"] is True
    assert r["colores"]["accent"] == "#233A5E"
    assert r["colores"]["gris"] == "#6E7684"


def test_api_template_colors_no_lanza():
    api = Api(None)
    for plantilla in (_Explosivo(), {"x": 1}, b"chinese", object()):
        r = api.template_colors(plantilla)
        assert isinstance(r, dict)
        assert isinstance(r.get("ok"), bool)
    r = api.template_colors(_Explosivo())
    assert r["ok"] is False
    assert isinstance(r["errores"], list) and r["errores"]


# ---------------------------------------------------------------------------
# docx_gen: palette precedence in the template-specific shapes
# ---------------------------------------------------------------------------

def _datos_docx():
    d = latex_gen.empty_data()
    d["cv_language"] = "es"
    d["name"] = "Ana García"
    d["summary"] = "Perfil de prueba."
    d["education"] = [{
        "degree_type": "bachelor", "degree": "Informática", "institution": "UM",
        "start": "2012", "end": "2016",
    }]
    d["work"] = [{"position": "Dev", "company": "Acme", "start": "2020", "end": "2024"}]
    d["blocks"] = [
        {"id": "linea", "type": "separator", "style": "line", "position": "end"},
        {"id": "banda", "type": "separator", "style": "band", "position": "end"},
    ]
    return d


def _exportar(datos, prefs_dict, tmp_path, nombre):
    """Export ``datos`` and reopen the .docx with python-docx."""
    ruta = tmp_path / nombre
    docx_gen.export(datos, prefs_dict, str(ruta))
    return Document(str(ruta))


# -- recursive readers ------------------------------------------------------

def _cell_blocks(cell):
    """Yield ``("paragraph"|"table", item)`` inside a table cell."""
    for child in cell._tc.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, cell)
        elif child.tag == qn("w:tbl"):
            table = Table(child, cell)
            yield "table", table
            for item in _table_blocks(table):
                yield item


def _table_blocks(table):
    for row in table.rows:
        for cell in row.cells:
            for item in _cell_blocks(cell):
                yield item


def _iter_blocks(doc):
    """Body and table blocks in document order (header table first)."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            yield "table", table
            for item in _table_blocks(table):
                yield item


def _iter_paragraphs(doc):
    for kind, item in _iter_blocks(doc):
        if kind == "paragraph":
            yield item


def _table_paragraphs(table):
    return [item for kind, item in _table_blocks(table) if kind == "paragraph"]


def _textos(doc):
    return [paragraph.text for paragraph in _iter_paragraphs(doc)]


def _ancestor_element(node, tag):
    element = node._p if isinstance(node, Paragraph) else node
    while element is not None:
        if element.tag == qn(tag):
            return element
        element = element.getparent()
    return None


def _ancestor_cell(paragraph):
    return _ancestor_element(paragraph, "w:tc")


# -- colours, fills and borders --------------------------------------------

def _shading_fill(element):
    if element is None:
        return None
    shading = element.find(qn("w:shd"))
    return shading.get(qn("w:fill")) if shading is not None else None


def _paragraph_fill(paragraph):
    return _shading_fill(paragraph._p.pPr)


def _cell_fill(cell):
    tc = getattr(cell, "_tc", cell)
    return _shading_fill(tc.find(qn("w:tcPr"))) if tc is not None else None


def _ancestor_cell_fill(paragraph):
    return _cell_fill(_ancestor_cell(paragraph))


def _borders(element, tag="w:pBdr"):
    if element is None:
        return {}
    container = element.find(qn(tag))
    if container is None:
        return {}
    out = {}
    for child in container:
        try:
            size = int(child.get(qn("w:sz")))
        except (TypeError, ValueError):
            size = None
        out[child.tag.rsplit("}", 1)[-1]] = {
            "val": child.get(qn("w:val")),
            "sz": size,
            "color": child.get(qn("w:color")),
        }
    return out


def _paragraph_borders(paragraph):
    return _borders(paragraph._p.pPr)


def _text_runs(paragraph):
    return [run for run in paragraph.runs if run.text]


def _run_color(run):
    color = run.font.color
    if color is None or color.type is None:
        return None
    return color.rgb


# -- header and heading locators -------------------------------------------

def _header_table(doc):
    assert doc.tables, "the header must be a table"
    return doc.tables[0]


def _header_name_paragraph(doc, name):
    for paragraph in _table_paragraphs(_header_table(doc)):
        if paragraph.text.strip() == name:
            return paragraph
    raise AssertionError("No name paragraph %r in the header table" % name)


def _header_name_runs(doc):
    return _text_runs(_header_name_paragraph(doc, _datos_docx()["name"]))


HEADING_MARKERS = "\u25cf\u25c6\u25aa\u25a0\u2022"


def _heading_text(text):
    return text.strip().lstrip(" \t" + HEADING_MARKERS)


def _find_heading(doc, title):
    """Heading dict: paragraph, table/cell and the shape depth (1 = table)."""
    for kind, item in _iter_blocks(doc):
        if kind == "paragraph":
            if _heading_text(item.text) == title:
                return {"paragraph": item, "table": None, "cell": None}
            continue
        if not item.rows:
            continue
        for cell in item.rows[0].cells:
            for paragraph in cell.paragraphs:
                if _heading_text(paragraph.text) == title:
                    return {"paragraph": paragraph, "table": item, "cell": cell}
    raise AssertionError("No section heading %r" % title)


def _heading_runs(heading):
    """Title runs only (the playful bullet glyph is not part of the title)."""
    runs = _text_runs(heading["paragraph"])
    if heading["table"] is None and runs and runs[0].text.strip() in set(HEADING_MARKERS):
        return runs[1:]
    return runs


def _heading_fill(heading):
    """Fill of the heading shape: plate/tab/bar cell, or paragraph shading."""
    if heading["table"] is not None:
        return _cell_fill(heading["cell"])
    return _paragraph_fill(heading["paragraph"])


def _paragraph_after(doc, paragraph):
    paragraphs = list(_iter_paragraphs(doc))
    for index, candidate in enumerate(paragraphs):
        if candidate._p is paragraph._p:
            return paragraphs[index + 1] if index + 1 < len(paragraphs) else None
    return None


def _heading_rule(doc, heading):
    """Bottom rule of a heading: its own paragraph, its rule cell or the next ¶."""
    if heading["table"] is not None:
        for cell in heading["table"].rows[0].cells:
            for paragraph in cell.paragraphs:
                border = _paragraph_borders(paragraph).get("bottom")
                if border:
                    return border
        heading_rule_paragraph = _paragraph_after(doc, heading["paragraph"])
    else:
        border = _paragraph_borders(heading["paragraph"]).get("bottom")
        if border:
            return border
        heading_rule_paragraph = _paragraph_after(doc, heading["paragraph"])
    if heading_rule_paragraph is None:
        return None
    return _paragraph_borders(heading_rule_paragraph).get("bottom")


def _header_rule_colors(doc):
    """Bottom-border colours of the rule paragraphs after the header table."""
    seen_table = False
    colors = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            if seen_table:
                break
            seen_table = True
            continue
        if not seen_table or child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc)
        if paragraph.text.strip():
            break
        border = _paragraph_borders(paragraph).get("bottom")
        if border:
            colors.append(border["color"])
    return colors


def _separadores(doc):
    """(line colours, band fills) of the trailing free-block separators.

    The two separator blocks are rendered at the "end" anchor, so they are the
    paragraphs after the last text of the document. A line separator is a thin
    bottom rule; a band separator is a shaded bar (or, for navy, another thin
    rule with no fill), and any secondary hairline of the profile is collected
    in the line list too.
    """
    paragraphs = list(_iter_paragraphs(doc))
    last_text = max(
        (i for i, paragraph in enumerate(paragraphs) if paragraph.text.strip()),
        default=-1,
    )
    lineas, bandas = [], []
    for paragraph in paragraphs[last_text + 1:]:
        border = _paragraph_borders(paragraph).get("bottom")
        if border:
            lineas.append(border["color"])
        fill = _paragraph_fill(paragraph)
        if fill:
            bandas.append(fill)
    return lineas, bandas


# -- tests ------------------------------------------------------------------

def test_docx_primary_y_text_con_paleta(tmp_path):
    """primary colours the header rule, heading titles/rules and separators."""
    prefs_colores = dict(PREFS_PRUEBA, colors={"primary": "#12AB34", "text": "#112233"})
    doc = _exportar(_datos_docx(), prefs_colores, tmp_path, "colores.docx")

    primario = RGBColor(0x12, 0xAB, 0x34)
    for titulo in ETIQUETAS_ES:
        heading = _find_heading(doc, titulo)
        # Classic heading: accent-coloured title over an accent rule, no fill.
        assert _heading_fill(heading) is None, titulo
        assert [_run_color(run) for run in _heading_runs(heading)] == [primario], titulo
        assert _heading_rule(doc, heading)["color"] == "12AB34", titulo

    # The header name takes the prefs text colour; the header rule the primary.
    assert [_run_color(run) for run in _header_name_runs(doc)] == [
        RGBColor(0x11, 0x22, 0x33)
    ]
    assert _header_rule_colors(doc) == ["12AB34"]

    lineas, bandas = _separadores(doc)
    assert lineas == ["12AB34"]
    assert bandas == ["12AB34"]

    assert doc.styles["Normal"].font.color.rgb == RGBColor(0x11, 0x22, 0x33)


def test_docx_sin_paleta_usa_el_acento_de_la_plantilla(tmp_path):
    doc = _exportar(_datos_docx(), PREFS_PRUEBA, tmp_path, "sin_colores.docx")

    # _datos_docx() does not pick a template: it falls back to classic (#233A5E).
    lineas, bandas = _separadores(doc)
    assert lineas == [docx_gen.COLOR_ACCENT]
    assert bandas == [docx_gen.COLOR_ACCENT]
    for titulo in ETIQUETAS_ES:
        heading = _find_heading(doc, titulo)
        assert _heading_fill(heading) is None, titulo
        assert [_run_color(run) for run in _heading_runs(heading)] == [
            RGBColor(0x23, 0x3A, 0x5E)
        ], titulo
        assert _heading_rule(doc, heading)["color"] == docx_gen.COLOR_ACCENT, titulo

    # The classic name sits on white in the text role (black when unset).
    assert [_run_color(run) for run in _header_name_runs(doc)] == [RGBColor(0, 0, 0)]
    assert _header_rule_colors(doc) == [docx_gen.COLOR_ACCENT]


def test_docx_paleta_invalida_cae_al_acento_de_plantilla(tmp_path):
    """An unusable prefs palette leaves every shape on the template accent."""
    prefs_colores = dict(PREFS_PRUEBA, colors={
        "primary": "red", "secondary": "#12", "gray": "  ", "text": "#GGGGGG",
    })
    doc = _exportar(_datos_docx(), prefs_colores, tmp_path, "invalida.docx")

    lineas, bandas = _separadores(doc)
    assert lineas == [docx_gen.COLOR_ACCENT]
    assert bandas == [docx_gen.COLOR_ACCENT]
    heading = _find_heading(doc, "Perfil")
    assert _heading_rule(doc, heading)["color"] == docx_gen.COLOR_ACCENT
    assert _header_rule_colors(doc) == [docx_gen.COLOR_ACCENT]
    assert doc.styles["Normal"].font.color.rgb is None


def test_docx_sword_usa_su_acento_en_cabecera_y_barras(tmp_path):
    d = _datos_docx()
    d["template"] = "sword"
    doc = _exportar(d, PREFS_PRUEBA, tmp_path, "sword.docx")

    # The sword header is a nameplate: navy banner cell with a white name.
    assert _ancestor_cell_fill(_header_name_paragraph(doc, d["name"])) == "1F346E"
    assert [_run_color(run) for run in _header_name_runs(doc)] == [
        RGBColor(0xFF, 0xFF, 0xFF)
    ]
    for titulo in ETIQUETAS_ES:
        heading = _find_heading(doc, titulo)
        # Accent plate behind the white title, accent2 full-width bar after it.
        assert _heading_fill(heading) == "1F346E", titulo
        assert [_run_color(run) for run in _heading_runs(heading)] == [
            RGBColor(0xFF, 0xFF, 0xFF)
        ], titulo
        cells = heading["table"].rows[0].cells
        assert _cell_fill(cells[2]) == "C4C8D4", titulo
        steps = [
            _paragraph_fill(paragraph)
            for paragraph in cells[1].paragraphs
            if _paragraph_fill(paragraph)
        ]
        assert steps and set(steps) == {"1F346E"}, titulo

    lineas, bandas = _separadores(doc)
    assert lineas == ["1F346E"]
    assert bandas == ["1F346E"]

    textos = _textos(doc)
    etiqueta_vieja = "个人简历 PERSONAL RESUME"
    assert all(etiqueta_vieja not in texto for texto in textos)


def test_docx_classic_nombre_negro_y_reglas_acento(tmp_path):
    """The classic name is black; the accent lives in its rules, not a fill."""
    d = _datos_docx()
    d["template"] = "classic"
    doc = _exportar(d, PREFS_PRUEBA, tmp_path, "classic.docx")

    assert _ancestor_cell_fill(_header_name_paragraph(doc, d["name"])) is None
    assert [_run_color(run) for run in _header_name_runs(doc)] == [RGBColor(0, 0, 0)]
    assert _header_rule_colors(doc) == ["233A5E"]

    heading = _find_heading(doc, "Perfil")
    assert _heading_fill(heading) is None
    assert [_run_color(run) for run in _heading_runs(heading)] == [
        RGBColor(0x23, 0x3A, 0x5E)
    ]
    assert _heading_rule(doc, heading)["color"] == "233A5E"


def test_docx_pref_primary_gana_a_la_plantilla(tmp_path):
    d = _datos_docx()
    d["template"] = "sword"
    prefs_colores = dict(PREFS_PRUEBA, colors={"primary": "#12AB34"})
    doc = _exportar(d, prefs_colores, tmp_path, "sword_primary.docx")

    # The palette reaches the sword-specific shapes: the banner and every
    # accent plate/blade, not a universal heading fill. The banner keeps its
    # white name (the LaTeX sword banner sets the name in white).
    assert _ancestor_cell_fill(_header_name_paragraph(doc, d["name"])) == "12AB34"
    assert [_run_color(run) for run in _header_name_runs(doc)] == [
        RGBColor(0xFF, 0xFF, 0xFF)
    ]

    heading = _find_heading(doc, "Perfil")
    assert _heading_fill(heading) == "12AB34"
    cells = heading["table"].rows[0].cells
    assert _cell_fill(cells[2]) == "C4C8D4"  # accent2 stays the template gray
    steps = [
        _paragraph_fill(paragraph)
        for paragraph in cells[1].paragraphs
        if _paragraph_fill(paragraph)
    ]
    assert steps and set(steps) == {"12AB34"}

    lineas, bandas = _separadores(doc)
    assert lineas == ["12AB34"]
    assert bandas == ["12AB34"]


def test_docx_acento_claro_usa_texto_contrastante(tmp_path):
    """A light accent still picks dark ink; the filled heading band is gone.

    The per-template shapes use fixed white-on-band names, so the light/dark
    switch survives in the shared :func:`docx_gen._contrast_color` helper; the
    assertion is kept at that level and the pref propagation is checked on the
    classic rules.
    """
    amarillo = RGBColor(0xF4, 0xB4, 0x40)
    assert docx_gen._contrast_color(amarillo) == RGBColor(0, 0, 0)
    assert docx_gen._contrast_color(RGBColor(0x23, 0x3A, 0x5E)) == RGBColor(0xFF, 0xFF, 0xFF)

    d = _datos_docx()
    d["template"] = "classic"
    prefs_colores = dict(PREFS_PRUEBA, colors={"primary": "#F4B440"})
    doc = _exportar(d, prefs_colores, tmp_path, "acento_claro.docx")

    heading = _find_heading(doc, "Perfil")
    assert _heading_fill(heading) is None
    assert _heading_rule(doc, heading)["color"] == "F4B440"
    assert _header_rule_colors(doc) == ["F4B440"]
    assert [_run_color(run) for run in _header_name_runs(doc)] == [RGBColor(0, 0, 0)]


def test_docx_creative_secundario_pisa_acento2_pero_no_acento3(tmp_path):
    """prefs secondary recolours acento2; the yellow acento3 is template-only."""
    d = _datos_docx()
    d["template"] = "creative"
    prefs_colores = dict(
        PREFS_PRUEBA,
        colors={"primary": "#12AB34", "secondary": "#00AA88", "gray": "#A1B2C3"},
    )
    doc = _exportar(d, prefs_colores, tmp_path, "creative.docx")

    heading = _find_heading(doc, "Perfil")
    assert _heading_rule(doc, heading)["color"] == "00AA88"
    segundo = _paragraph_after(doc, heading["paragraph"])
    assert segundo is not None
    assert _paragraph_borders(segundo)["bottom"]["color"] == "F4B440"

    # The palette reaches the creative band (the template-specific shape).
    assert _ancestor_cell_fill(_header_name_paragraph(doc, d["name"])) == "12AB34"
    assert [_run_color(run) for run in _heading_runs(heading)] == [
        RGBColor(0x12, 0xAB, 0x34)
    ]


def test_docx_gray_pisa_el_gris_de_texto_no_la_barra(tmp_path):
    """prefs gray overrides the readable text gray, never the light bar gray."""
    d = _datos_docx()
    d["template"] = "chinese"
    prefs_colores = dict(PREFS_PRUEBA, colors={"gray": "#A1B2C3", "text": "#112233"})
    doc = _exportar(d, prefs_colores, tmp_path, "chinese_gray.docx")

    grado = next(
        paragraph for paragraph in _iter_paragraphs(doc)
        if "Grado Universitario" in paragraph.text
    )
    assert [_run_color(run) for run in _text_runs(grado)] == [
        RGBColor(0xA1, 0xB2, 0xC3)
    ]
    heading = _find_heading(doc, "Perfil")
    assert _heading_fill(heading) == "1F347D"          # template accent plate
    assert _cell_fill(heading["table"].rows[0].cells[1]) == "C4C8D4"  # raw gris
    assert doc.styles["Normal"].font.color.rgb == RGBColor(0x11, 0x22, 0x33)


def test_docx_datos_y_prefs_hostiles_no_lanzan(tmp_path):
    datos_hostiles = [_DictExplosivo(), _Explosivo(), None, {}, {
        "blocks": _Explosivo(),
        "education": _Explosivo(),
        "summary": _Explosivo(),
        "work": [{"position": _Explosivo(), "achievements": _Explosivo()}],
        "skills": [{"category": _Explosivo(), "values": _Explosivo()}],
        "socials": [{"network": _Explosivo(), "value": _Explosivo()}],
    }]
    for i, datos in enumerate(datos_hostiles):
        ruta = tmp_path / ("hostil_%d.docx" % i)
        docx_gen.export(datos, _DictExplosivo(paper="a4"), str(ruta))
        assert ruta.is_file()


# ---------------------------------------------------------------------------
# Compilacion real de un .tex con paleta
# ---------------------------------------------------------------------------

@pytest.fixture()
def prefs_tmp(tmp_path, monkeypatch):
    """Aisla prefs (dirs y cache) del equipo del usuario."""
    app = tmp_path / "app"
    monkeypatch.setattr(prefs, "APP_DIR", str(app))
    monkeypatch.setattr(prefs, "CACHE_DIR", str(app / "cache"))
    monkeypatch.setattr(prefs, "BUILD_DIR", str(app / "cache" / "build"))
    monkeypatch.setattr(prefs, "PREFS_FILE", str(app / "prefs.json"))
    monkeypatch.setattr(prefs, "_CACHE", None)
    return tmp_path


@pytest.mark.latex
@pytest.mark.skipif(MOTOR_PDF is None, reason="Sin motor LaTeX (pdflatex) instalado")
def test_compila_pdf_con_paleta(prefs_tmp):
    tex = latex_gen.build_tex(_datos_prueba(), _prefs_con(COLORES_VALIDOS))
    resultado = latex_engine.compile(tex, None, None)
    assert resultado["ok"], resultado["errores"]
    assert resultado["pdf"] and os.path.isfile(resultado["pdf"])
